#!/usr/bin/env python3
"""GEOPOD v2 - Importer data Podes Excel (.xlsx) -> JSONB potensi_desa.

Membaca workbook Podes (sheet "ALL" + sheet "dict" + sheet section I, III,
V, VI, VII, VIII, IX, X, XI, XII), lalu meng-insert/update
`kelurahan_attribute.potensi_desa` diorganisasi:

    potensi_desa = {
      "meta":     {"iddesa": ..., "tahun": ..., "nama_prov": ..., ...},
      "sections": { "<section>": {"label": "...", "fields": [
                        {"k": "r101", "l": "Kode desa/kelurahan", "v": "..."}, ... ]}, ... }
    }

Nilai kosong dimasukkan sebagai "–" supaya setiap kelurahan menampilkan
seluruh rincian kolom (seragam). Butuh geometri sudah ter-seed
(kelurahan_geom) karena ada FOREIGN KEY.

Mode revisi label: data jawaban boleh datang dari file terpisah (mis.
`datageopodlabel.xlsx` yang isi `v` berupa label kuesioner) via `--src`,
sedangkan metadata (sheet "dict" untuk label kolom + sheet section untuk
peta kode -> section) tetap diambil dari `--ref` (default = `--src`).
Sheet data dideteksi otomatis: "ALL", atau sheet apa pun yang baris
pertamanya memuat kolom `iddesa`.

Penggunaan:
    python import_podes_xlsx.py --src ./podes_seed/datageopodlabel.xlsx \
        --ref ./podes_seed/datageopod.xlsx
Env: DATABASE_URL / GEOPOD_DB_URL
"""
import argparse
import json
import os
import re
import sys
from pathlib import Path

SECTION_NAMES = {
    "i.":  ("keterangan_tempat", "Informasi Umum & Demografi (Keterangan Tempat)"),
    "iii.": ("umum_kelurahan", "Keterangan Umum Kelurahan"),
    "v.":  ("perumahan_lingkungan", "Perumahan dan Lingkungan Hidup"),
    "vi.": ("bencana_mitigasi", "Bencana Alam dan Mitigasi"),
    "vii.": ("pendidikan_kesehatan", "Pendidikan dan Kesehatan"),
    "viii.": ("angkutan_komunikasi", "Angkutan, Komunikasi dan Informasi"),
    "ix.": ("ekonomi", "Ekonomi"),
    "x.":  ("keamanan", "Keamanan"),
    "xi.": ("keuangan_aset", "Keuangan dan Aset Desa"),
    "xii.": ("aparatur", "Keterangan Aparatur Pemerintahan"),
}


def norm_code(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return str(int(v)) if float(v).is_integer() else str(v)
    s = str(v).strip()
    if s and re.fullmatch(r"\d+\.0", s):
        s = s[:-2]
    return s or None


def load_workbook(path):
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        sys.exit(f"[podes] openpyxl tidak terpasang di container: {e}")
    return load_workbook(str(path), read_only=True, data_only=True)


def first_row_values(ws):
    """Nilai baris pertama sebuah sheet (mengisi hingga kolom terakhir)."""
    row = next(ws.iter_rows(min_row=1, max_row=1), None)
    vals = []
    if row:
        for cell in row:
            vals.append(cell.value)
    return vals


def build_code_label_map(dict_ws, all_headers):
    """Bangun peta kolom -> label dari sheet 'dict' dengan deteksi arah otomatis."""
    map_ = {}
    if dict_ws is None:
        return map_
    header_set = {h for h in all_headers if h}
    for row in dict_ws.iter_rows(values_only=True):
        a, b = (row[0] if len(row) > 0 else None), (row[1] if len(row) > 1 else None)
        na, nb = norm_code(a), norm_code(b)
        if na in header_set and nb not in header_set:
            map_[na] = nb
        elif nb in header_set and na not in header_set:
            map_[nb] = na
    return map_


def detail_label(sec):
    for _slug, label in SECTION_NAMES.values():
        if _slug == sec:
            return label
    return sec


def main():
    ap = argparse.ArgumentParser(description="Importer Podes Excel -> JSONB")
    ap.add_argument("--src", required=True, help="Path ke file data Podes (.xlsx), isi nilai jawaban (boleh berlabel).")
    ap.add_argument("--ref", default=None, help="Path file metadata (dict + sheet section). Default: --src.")
    ap.add_argument("--tahun", default="2025")
    ap.add_argument("--dry-run", action="store_true",
                    help="Parsing saja (tanpa koneksi DB), tampilkan ringkasan.")
    args = ap.parse_args()

    src = Path(args.src)
    if not src.exists():
        sys.exit(f"[podes] File tidak ada: {src}")

    wb = load_workbook(src)
    ref_path = Path(args.ref) if args.ref else src
    wb_ref = wb if ref_path == src else load_workbook(ref_path)

    all_ws = None
    for ws in wb.worksheets:
        if (ws.title or "").strip().lower() == "all":
            all_ws = ws
            break
    if all_ws is None:
        for ws in wb.worksheets:
            if any(norm_code(c) == "iddesa" for c in first_row_values(ws)):
                all_ws = ws
                break
    if all_ws is None:
        sys.exit("[podes] Sheet data (ALL atau sheet berisi 'iddesa') tidak ditemukan di --src.")

    sec_map = {}  # code -> section_key
    sec_title = {}  # section_key -> judul asli sheet (ws.title)
    dict_ws = None
    for ws in wb_ref.worksheets:
        name = (ws.title or "").strip()
        if name.lower() == "dict":
            dict_ws = ws
            continue
        slug = re.findall(r"^([a-z]+)\.?", name.lower()) or []
        key = slug[0] + "." if slug else ""
        sec = SECTION_NAMES.get(key)
        if not sec:
            continue
        sec_title[sec[0]] = name
        for code in first_row_values(ws):
            c = norm_code(code)
            if c:
                sec_map.setdefault(c, sec[0])

    headers = [norm_code(c) for c in first_row_values(all_ws)]
    label_map = build_code_label_map(dict_ws, headers)
    # Kolom identitas yang ada di ALL tetapi tidak punya sheet section -> meta
    meta_codes = {"iddesa", "nama_prov", "nama_kab", "nama_kec", "nama_desa"}
    for c in headers:
        if c and c not in meta_codes and c.startswith("r") and c not in sec_map:
            # default ke section keterangan_tempat untuk kode r1xx yang tak terpetakan
            if c.startswith("r1"):
                sec_map[c] = "keterangan_tempat"

    try:
        from psycopg2 import connect
    except ImportError as e:
        if not args.dry_run:
            sys.exit(f"[podes] psycopg2 tidak terpasang di container: {e}")

    url = os.environ.get("DATABASE_URL") or os.environ.get("GEOPOD_DB_URL")
    if args.dry_run:
        conn = None
    else:
        if not url:
            sys.exit("[podes] Set DATABASE_URL / GEOPOD_DB_URL.")
        conn = connect(url)

    ok = skipped = 0
    sample = None
    for row in all_ws.iter_rows(min_row=2, values_only=True):
        cells = list(row or [])
        if not any(v is not None and str(v).strip() for v in cells):
            continue
        data = {}
        for i, h in enumerate(headers):
            if not h:
                continue
            v = cells[i] if i < len(cells) else None
            if v is None or str(v).strip() == "":
                continue
            data[h] = str(v).strip()

        iddesa = norm_code(data.get("iddesa"))
        nama_desa = data.get("nama_desa")
        if not iddesa:
            skipped += 1
            continue

        sections = {}
        for i, h in enumerate(headers):
            if not h or h in meta_codes or h == "iddesa":
                continue
            v = cells[i] if i < len(cells) else None
            sv = "" if v is None else str(v).strip()
            sec = sec_map.get(h, "umum_kelurahan")
            group = sections.setdefault(sec, [])
            label = label_map.get(h, h)
            group.append({"k": h, "l": label, "v": sv if sv else "–"})

        potensi = {
            "meta": {
                "tahun": norm_code(args.tahun),
                "iddesa": iddesa,
                "nama_prov": data.get("nama_prov"),
                "nama_kab": data.get("nama_kab"),
                "nama_kec": data.get("nama_kec"),
                "nama_desa": nama_desa,
            },
            "sections": {
                sec: {
                    "label": sec_title.get(sec, detail_label(sec)),
                    "fields": fields,
                }
                for sec, fields in sorted(sections.items())
            },
        }

        if args.dry_run:
            ok += 1
            if sample is None:
                sample = potensi
            continue

        with conn.cursor() as cur:
            cur.execute(
                "SELECT 1 FROM kelurahan_geom WHERE kode_kelurahan = %s", (iddesa,)
            )
            if cur.fetchone() is None:
                sys.stderr.write(f"[podes] geometri {iddesa} belum ada -> dilewati.\n")
                skipped += 1
                continue
            cur.execute(
                """
                INSERT INTO kelurahan_attribute (kode_kelurahan, nama_kelurahan,
                                                 nama_kecamatan, potensi_desa)
                VALUES (%s, %s, %s, %s)
                ON CONFLICT (kode_kelurahan)
                DO UPDATE SET nama_kelurahan = EXCLUDED.nama_kelurahan,
                              nama_kecamatan = EXCLUDED.nama_kecamatan,
                              potensi_desa = EXCLUDED.potensi_desa
                """,
                (iddesa, nama_desa, data.get("nama_kec"), json.dumps(potensi, ensure_ascii=False)),
            )
        ok += 1

    if args.dry_run:
        print(f"[podes][dry-run] {ok} kelurahan ter-parse (sample: {sample and sample['meta'].get('nama_desa')}).")
        if sample:
            print(json.dumps(sample, ensure_ascii=False, indent=2)[:4000])
        return

    conn.commit()
    conn.close()
    print(f"[podes] Selesai: {ok} kelurahan di-import, {skipped} dilewati.")


if __name__ == "__main__":
    main()